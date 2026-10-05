"""Step 1 of the pipeline: text extraction for PDF, DOCX, TXT and CSV."""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import pandas as pd
from docx import Document
from pypdf import PdfReader


@dataclass
class Segment:
    """A piece of raw text plus where it came from (page number, row range...)."""
    text: str
    location: str = ""


@dataclass
class LoadedDocument:
    name: str
    file_type: str
    segments: list[Segment] = field(default_factory=list)
    dataframe: pd.DataFrame | None = None  # only for CSV

    @property
    def full_text(self) -> str:
        return "\n\n".join(s.text for s in self.segments)


def _load_pdf(data: bytes) -> list[Segment]:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    segments = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segments.append(Segment(text, f"page {i}"))
    return segments


def _load_docx(data: bytes) -> list[Segment]:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [Segment(text, "document")] if text else []


def _load_txt(data: bytes) -> list[Segment]:
    text = ""
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(enc).strip()
            break
        except UnicodeDecodeError:
            continue
    return [Segment(text, "document")] if text else []


def _load_csv(data: bytes) -> tuple[list[Segment], pd.DataFrame]:
    try:
        df = pd.read_csv(io.BytesIO(data))
    except UnicodeDecodeError:
        df = pd.read_csv(io.BytesIO(data), encoding="latin-1")
    if df.empty:
        return [], df

    segments = [Segment(
        f"CSV file with {len(df)} rows and {len(df.columns)} columns.\n"
        f"Columns: {', '.join(map(str, df.columns))}\n\n"
        f"Summary statistics:\n{df.describe(include='all').to_string()}",
        "overview",
    )]
    # Rows are grouped so each segment stays a sensible size.
    step = 25
    for start in range(0, len(df), step):
        block = df.iloc[start:start + step]
        rows = "\n".join(
            "; ".join(f"{c}={v}" for c, v in r.items()) for _, r in block.iterrows()
        )
        segments.append(Segment(rows, f"rows {start + 1}-{start + len(block)}"))
    return segments, df


def load_document(name: str, data: bytes) -> LoadedDocument:
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "pdf":
        doc = LoadedDocument(name, ext, _load_pdf(data))
    elif ext == "docx":
        doc = LoadedDocument(name, ext, _load_docx(data))
    elif ext == "txt":
        doc = LoadedDocument(name, ext, _load_txt(data))
    elif ext == "csv":
        segments, df = _load_csv(data)
        doc = LoadedDocument(name, ext, segments, df)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    if not doc.segments:
        raise ValueError(
            "No text could be extracted. If this is a scanned PDF, it needs OCR first."
        )
    return doc
